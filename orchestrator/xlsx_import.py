"""Parse external SAST spreadsheet alerts into canonical warnings."""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Mapping

CONTRACTS_PYTHON = Path(__file__).resolve().parents[1] / "contracts" / "python"
if str(CONTRACTS_PYTHON) not in sys.path:
    sys.path.insert(0, str(CONTRACTS_PYTHON))

from warning_contract import new_warning, recompute_score, require_warning  # noqa: E402

REQUIRED_COLUMNS = {
    "规范": "norm",
    "缺陷描述": "snippet",
    "文件": "file",
    "代码行": "line",
    "规则名": "rule_name",
}

DEFAULT_PRODUCER = "tabular-sast"
DEFAULT_INIT_MODEL = "xlsx-init"


def normalize_path(path: str) -> str:
    normalized = str(path or "").replace("\\", "/").strip()
    while normalized.startswith("./"):
        normalized = normalized[2:]
    return normalized


def same_or_suffix_path(left: str, right: str) -> bool:
    left_n = normalize_path(left).lstrip("/")
    right_n = normalize_path(right).lstrip("/")
    return bool(left_n and right_n) and (
        left_n == right_n
        or left_n.endswith("/" + right_n)
        or right_n.endswith("/" + left_n)
    )


def _source_suffix_prefixes(source: str) -> list[str]:
    """Return trailing path prefixes of ``source`` (longest first), e.g. a/b/c -> a/b/c, b/c, c."""
    parts = [part for part in normalize_path(source).split("/") if part]
    return ["/".join(parts[index:]) for index in range(len(parts))]


def path_matches_source(file_path: str, source_dir: Path) -> bool:
    """True when the alert file belongs to ``source_dir`` by path heuristics.

    Relative spreadsheet paths must not be treated as matches merely because
    ``source_dir / relative`` always resolves under ``source_dir``; that used to
    import every row. Prefer real files under ``source_dir``, then path-prefix
    / leaf heuristics against ``workflow.ini`` ``source_dir``.
    """
    file_name = normalize_path(file_path)
    if not file_name:
        return False
    try:
        source = normalize_path(str(source_dir.resolve()))
    except OSError:
        source = normalize_path(str(source_dir))
    source_leaf = Path(source).name

    try:
        candidate = Path(file_path).expanduser()
        if candidate.is_absolute():
            resolved = normalize_path(str(candidate.resolve()))
            if resolved == source or resolved.startswith(source + "/"):
                return True
        else:
            # Only trust join when the file actually exists under source_dir.
            joined = (source_dir / candidate).resolve()
            resolved = normalize_path(str(joined))
            under_source = resolved == source or resolved.startswith(source + "/")
            if under_source and joined.exists():
                return True
    except OSError:
        pass

    if same_or_suffix_path(file_name, source) or same_or_suffix_path(file_name, source_leaf):
        return True

    # e.g. source_dir=.../drivers/ub and file=drivers/ub/mem/foo.c
    padded = f"/{file_name}/"
    for prefix in _source_suffix_prefixes(source):
        if file_name == prefix or file_name.startswith(prefix + "/"):
            return True
        if f"/{prefix}/" in padded:
            return True

    parts = [part for part in file_name.split("/") if part]
    if source_leaf and source_leaf in parts:
        return True
    return bool(source_leaf) and file_name.startswith(source_leaf + "/")


def alert_file_location(document: Mapping[str, Any]) -> str | None:
    content = document.get("content")
    if not isinstance(content, Mapping):
        return None
    if document.get("type") == "tabular":
        location = content.get("location")
        if isinstance(location, Mapping):
            return str(location.get("file") or "") or None
    path = content.get("path")
    if isinstance(path, list) and path:
        node = path[0]
        if isinstance(node, Mapping):
            loc = node.get("location")
            if isinstance(loc, Mapping) and loc.get("file"):
                return str(loc["file"])
    allocation = content.get("allocation")
    if isinstance(allocation, Mapping):
        loc = allocation.get("location")
        if isinstance(loc, Mapping) and loc.get("file"):
            return str(loc["file"])
    access = content.get("access")
    if isinstance(access, Mapping):
        loc = access.get("location")
        if isinstance(loc, Mapping) and loc.get("file"):
            return str(loc["file"])
    return None


def path_filter_document(document: Mapping[str, Any], source_dir: Path) -> bool:
    file_name = alert_file_location(document)
    if not file_name:
        return False
    return path_matches_source(file_name, source_dir)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_map(row: tuple[Any, ...]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, cell in enumerate(row):
        name = _cell_text(cell)
        if name in REQUIRED_COLUMNS and name not in mapping:
            mapping[name] = index
    missing = sorted(REQUIRED_COLUMNS.keys() - mapping.keys())
    if missing:
        raise ValueError("xlsx missing required column(s): " + ", ".join(missing))
    return mapping


def _parse_line(value: Any) -> int:
    text = _cell_text(value)
    if not text:
        raise ValueError("代码行 is empty")
    try:
        line = int(float(text))
    except ValueError as error:
        raise ValueError(f"invalid 代码行: {value!r}") from error
    if line < 1:
        raise ValueError(f"代码行 must be >= 1: {line}")
    return line


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ValueError(
            "openpyxl is required to import xlsx alerts; pip install openpyxl"
        ) from error
    if not path.is_file():
        raise ValueError(f"xlsx file does not exist: {path}")
    # read_only trusts sheet dimension; some exporters set it to one row and
    # drop the rest. Use normal mode so max_row reflects actual cells.
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration as error:
            raise ValueError("xlsx sheet is empty") from error
        columns = _header_map(tuple(header or ()))
        rows: list[dict[str, Any]] = []
        for row_number, raw in enumerate(rows_iter, start=2):
            values = tuple(raw or ())
            if not any(_cell_text(item) for item in values):
                continue
            try:
                payload = {
                    "norm": _cell_text(values[columns["规范"]]),
                    "snippet": _cell_text(values[columns["缺陷描述"]]),
                    "file": normalize_path(_cell_text(values[columns["文件"]])),
                    "line": _parse_line(values[columns["代码行"]]),
                    "rule_name": _cell_text(values[columns["规则名"]]),
                    "row_number": row_number,
                }
            except (IndexError, ValueError) as error:
                raise ValueError(f"xlsx row {row_number}: {error}") from error
            for key in ("norm", "snippet", "file", "rule_name"):
                if not payload[key]:
                    raise ValueError(f"xlsx row {row_number}: {key} is empty")
            rows.append(payload)
        return rows
    finally:
        workbook.close()


def build_tabular_warning(
    *,
    producer: str,
    norm: str,
    snippet: str,
    file_name: str,
    line: int,
    rule_name: str,
    initial_weight: float = 0.5,
    init_model: str = DEFAULT_INIT_MODEL,
) -> dict[str, Any]:
    weight = float(initial_weight)
    if not 0.0 <= weight <= 1.0:
        raise ValueError("initial_weight must be between 0 and 1")
    producer_text = producer.strip()
    if not producer_text:
        raise ValueError("producer must not be empty")
    content = {
        "location": {"file": normalize_path(file_name), "line": int(line)},
        "snippet": str(snippet).strip(),
        "norm": str(norm).strip(),
        "rule_name": str(rule_name).strip(),
    }
    warning = new_warning(producer_text, "tabular", content)
    warning["graph_ids"] = []
    warning["active_learning"] = {"weight": weight, "model": init_model}
    recompute_score(warning)
    require_warning(warning)
    return warning
